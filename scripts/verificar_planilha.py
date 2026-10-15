"""Etapa 7: confere as fórmulas da planilha com um motor de cálculo independente.

    python scripts/verificar_planilha.py [outro_municipio]

1. Recalcula todas as fórmulas do arquivo com a biblioteca `formulas` (um interpretador de fórmulas do
   Excel escrito em Python, que não compartilha nada com o xlsxwriter nem com este repositório).
2. Compara cada resultado com o valor que o script planilha.py gravou em cache.
3. Opcional: troca o município da célula Painel!C4 por outro, recalcula e confere o resultado contra os
   CSV tratados (prova que a seleção funciona para quem não é o município de exemplo).
Sai com código 1 se alguma fórmula der erro inesperado ou divergir do cache.
"""
import os
import sys

import formulas
import openpyxl

import comum
from planilha import SAIDA

ABAS_COM_FORMULA = ['Painel', 'Tabela do município', 'Previsão']


def calcular(caminho, troca=None):
    """Devolve {(aba, célula): valor} de todas as células de fórmula, recalculadas."""
    if troca:
        wb = openpyxl.load_workbook(caminho)
        wb['Painel']['C4'] = troca
        caminho = os.path.join(os.path.dirname(caminho), '_verificacao.xlsx')
        wb.save(caminho)
    modelo = formulas.ExcelModel().loads(caminho).finish()
    sol = modelo.calculate()
    base = os.path.basename(caminho).upper()
    saida = {}
    for chave, rng in sol.items():
        c = str(chave)
        if '!' not in c or ':' in c.split('!')[1]:
            continue
        aba, celula = c.split('!')
        if f'[{base}]' not in aba.upper():
            continue
        nome = aba.split(']')[1].strip("'")
        try:
            valor = rng.value[0, 0]
        except Exception:
            continue
        saida[(nome.upper(), celula.replace('$', ''))] = valor
    if troca:
        os.remove(caminho)
    return saida


def main():
    wb_f = openpyxl.load_workbook(SAIDA)                        # fórmulas
    wb_c = openpyxl.load_workbook(SAIDA, data_only=True)        # valores em cache
    calc = calcular(SAIDA)
    divergentes, conferidas, erros = [], 0, []
    for aba in ABAS_COM_FORMULA:
        for linha in wb_f[aba].iter_rows():
            for cel in linha:
                if isinstance(cel.value, str) and cel.value.startswith('='):
                    cache = wb_c[aba][cel.coordinate].value
                    novo = calc.get((aba.upper(), cel.coordinate))
                    conferidas += 1
                    if novo is None:
                        erros.append((aba, cel.coordinate, 'não calculada'))
                        continue
                    texto = str(novo)
                    if texto.startswith('#') and cache != '#N/A':
                        erros.append((aba, cel.coordinate, texto))
                        continue
                    if isinstance(cache, (int, float)) and not isinstance(novo, str):
                        if abs(float(novo) - float(cache)) > 1e-6 * max(1.0, abs(float(cache))):
                            divergentes.append((aba, cel.coordinate, cache, novo))
                    elif str(cache) != str(novo) and not (cache in (None, '') and str(novo) in ('', 'nan')):
                        # NA() devolve erro no motor e "#N/A" no cache: equivalentes
                        if not (str(cache) == '#N/A' and texto.startswith('#N/A')):
                            divergentes.append((aba, cel.coordinate, cache, novo))
    print(f'{conferidas} fórmulas conferidas nas abas {ABAS_COM_FORMULA}')
    print(f'erros inesperados: {len(erros)}; divergências do cache: {len(divergentes)}')
    for e in erros[:20]:
        print('  erro', e)
    for d in divergentes[:20]:
        print('  diverge', d)

    # prova de que a seleção funciona: outro município, outra faixa
    alvo = sys.argv[1] if len(sys.argv) > 1 else None
    if alvo:
        calc2 = calcular(SAIDA, troca=alvo)
        dim = {r['municipio']: r for r in comum.ler_csv('dim_municipio.csv')}
        cod = int(dim[alvo]['cod_ibge'])
        val = {(int(r['cod_ibge']), int(r['ano']), r['indicador_id']): float(r['valor']) for r in comum.ler_csv('indicadores_longo.csv')}
        ok = True
        # Tabela do município: linha 10 (receita por habitante), colunas C..M = 2015..2025
        for j, ano in enumerate(comum.ANOS):
            esperado = val.get((cod, ano, 'receita_pc'))
            obtido = calc2.get(('TABELA DO MUNICÍPIO', f'{chr(ord("C") + j)}10'))
            if esperado is None:
                bom = str(obtido) in ('n.d.',)
            else:
                bom = obtido is not None and not isinstance(obtido, str) and abs(float(obtido) - esperado) < 1e-6 * esperado
            if not bom:
                ok = False
                print('  FALHA', alvo, ano, 'esperado', esperado, 'obtido', obtido)
        print(f'seleção de "{alvo}": ' + ('receita por habitante confere nos 11 anos' if ok else 'DIVERGE'))
        if not ok:
            sys.exit(1)
    if erros or divergentes:
        sys.exit(1)


if __name__ == '__main__':
    main()
